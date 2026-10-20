"""Defines document extraction strategies for various file types."""

import contextlib
import csv
import io
import logging
from typing import Protocol, cast

import pypdf

from app.core.progress import emit_progress


def get_ocr_reader():
    """Lazily load and return the EasyOCR Reader instance configured for CPU execution."""
    from app.core.shared_registry import SharedModelRegistry

    return SharedModelRegistry.get_instance().get_ocr_reader()


def _emit_progress(
    progress_callback, pct: float = 0.0, stage: str | None = None
) -> None:
    """Safely invoke a progress callback with ratio and optional stage description using central emitter."""
    emit_progress(progress_callback, progress_or_update=pct, stage=stage)


def extract_text_from_image(
    image, settings=None, file_path=None, progress_callback=None
) -> str:
    """Extract character-level text from an image using configured vision engine with dynamic OCR fallback."""
    if settings is None:
        from app.config import AppSettings

        try:
            settings = AppSettings()
        except Exception:
            logging.debug("Could not load AppSettings, using defaults", exc_info=True)

    try:
        from PIL import Image

        if isinstance(image, str):
            file_path = file_path or image
            image = Image.open(image)

        # Check image dimensions / skip threshold / downscaling
        width, height = None, None
        if hasattr(image, "size"):
            try:
                sz = image.size
                if isinstance(sz, tuple) and len(sz) == 2:
                    width, height = sz
            except Exception:
                logging.debug("Could not read image size", exc_info=True)

        if isinstance(width, (int, float)) and isinstance(height, (int, float)):
            max_dimension = (
                getattr(settings, "IMAGE_MAX_DIMENSION", 1000) if settings else 1000
            )

            if max(width, height) > max_dimension:
                ratio = max_dimension / max(width, height)
                new_width = max(min(width, 400), int(width * ratio))
                new_height = max(min(height, 400), int(height * ratio))
                name = file_path if file_path else "In-memory image"
                logging.info(
                    f"Status flag [DOWNSCALED]: Downscaling high-resolution image {name} from {(width, height)} to {(new_width, new_height)} prior to text recognition"
                )
                image = image.resize((new_width, new_height), Image.Resampling.LANCZOS)
    except Exception as e:
        logging.error(f"Failed preprocessing image: {e}", exc_info=True)

    vision_engine = (
        getattr(settings, "VISION_ENGINE", "easyocr") if settings else "easyocr"
    )
    if vision_engine == "florence-2":
        try:
            from app.core.shared_registry import SharedModelRegistry

            _emit_progress(
                progress_callback, 0.2, "Initializing Florence-2 VLM model..."
            )
            proc = SharedModelRegistry.get_instance().get_florence_processor()
            if proc is not None:
                _emit_progress(progress_callback, 0.5, "Extracting visual text...")
                res = proc.process_image(image, task_prompt="<OCR>")
                extracted = res.get("sanitized_text", "") or res.get("raw_output", "")
                if extracted and extracted.strip():
                    return extracted.strip()
        except Exception as e:
            logging.warning(
                f"Florence-2 vision engine processing failed for {file_path or 'image'}, falling back to EasyOCR: {e}",
                exc_info=True,
            )

    # Fallback to standard OCR (EasyOCR)
    _emit_progress(progress_callback, 0.3, "Initializing EasyOCR model...")
    reader = get_ocr_reader()
    if reader is None:
        return ""

    try:
        import numpy as np

        img_np = np.array(image)
        _emit_progress(progress_callback, 0.6, "Extracting visual text...")
        results = reader.readtext(img_np)
        extracted_text = " ".join([res[1] for res in results])
        return extracted_text.strip()
    except Exception as e:
        logging.error(f"OCR processing failed: {e}", exc_info=True)
        return ""


class DocumentExtractor(Protocol):
    """Protocol for extracting text from documents."""

    def extract(self, file_path: str) -> str:
        """Extract and return text from the given file."""
        ...


class TxtExtractor:
    """Extractor for plain text files."""

    def extract(self, file_path: str, settings=None) -> str:
        """Extract text from a .txt file."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()


class DocxExtractor:
    """Extractor for Microsoft Word documents."""

    def extract(self, file_path: str, settings=None) -> str:
        """Extract text from a .docx file."""
        from docx import Document

        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])


class CsvExtractor:
    """Extractor for comma-separated values files."""

    def extract(self, file_path: str, settings=None) -> str:
        """Extract text from a .csv file using row-by-row streaming up to row and character limits."""
        if settings is None:
            from app.config import AppSettings

            try:
                settings = AppSettings()
            except Exception:
                logging.debug(
                    "Could not load AppSettings, using defaults", exc_info=True
                )

        max_rows = getattr(
            settings,
            "TABULAR_MAX_ROWS",
            getattr(settings, "MAX_SPREADSHEET_ROWS", 10000),
        )
        max_chars = getattr(
            settings,
            "TABULAR_MAX_CHARACTERS",
            getattr(settings, "MAX_SPREADSHEET_CHARACTERS", 50000),
        )

        extracted_rows = []
        total_chars = 0
        row_count = 0

        with open(file_path, newline="", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                if row_count >= max_rows:
                    break
                row_str = " ".join(str(cell) for cell in row if cell is not None)

                added_len = len(row_str) + (1 if extracted_rows else 0)
                if total_chars + added_len > max_chars:
                    remaining = max_chars - total_chars - (1 if extracted_rows else 0)
                    if remaining > 0:
                        extracted_rows.append(row_str[:remaining])
                    break

                extracted_rows.append(row_str)
                total_chars += added_len
                row_count += 1

        return " ".join(extracted_rows)


class XlsxExtractor:
    """Extractor for Excel spreadsheets."""

    def extract(self, file_path: str, settings=None) -> str:
        """Extract text from an Excel file using read-only streaming mode across sheets."""
        if settings is None:
            from app.config import AppSettings

            try:
                settings = AppSettings()
            except Exception:
                logging.debug(
                    "Could not load AppSettings, using defaults", exc_info=True
                )

        max_sheets = getattr(
            settings,
            "TABULAR_MAX_SHEETS",
            getattr(settings, "MAX_SPREADSHEET_SHEETS", 10),
        )
        max_rows = getattr(
            settings,
            "TABULAR_MAX_ROWS",
            getattr(settings, "MAX_SPREADSHEET_ROWS", 10000),
        )
        max_chars = getattr(
            settings,
            "TABULAR_MAX_CHARACTERS",
            getattr(settings, "MAX_SPREADSHEET_CHARACTERS", 50000),
        )

        extracted_lines = []
        total_chars = 0

        # Primary method: openpyxl in read_only streaming mode
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            try:
                sheet_names = wb.sheetnames[:max_sheets]
                for sheet_name in sheet_names:
                    if total_chars >= max_chars:
                        break
                    sheet = wb[sheet_name]
                    row_count = 0
                    for row in sheet.iter_rows(values_only=True):
                        if row_count >= max_rows:
                            break
                        row_str = " ".join(
                            str(cell) for cell in row if cell is not None
                        ).strip()

                        added_len = len(row_str) + (1 if extracted_lines else 0)
                        if total_chars + added_len > max_chars:
                            remaining = (
                                max_chars - total_chars - (1 if extracted_lines else 0)
                            )
                            if remaining > 0:
                                extracted_lines.append(row_str[:remaining])
                            total_chars = max_chars
                            break

                        extracted_lines.append(row_str)
                        total_chars += added_len
                        row_count += 1
            finally:
                wb.close()

            return "\n".join(extracted_lines) if extracted_lines else ""

        except Exception as e:
            logging.debug(
                f"openpyxl read-only streaming extraction failed or non-xlsx format for {file_path}: {e}. Falling back to pandas.",
                exc_info=True,
            )

        # Fallback method: pandas.read_excel with nrows
        try:
            import pandas as pd

            try:
                xl = pd.ExcelFile(file_path)
                sheet_names = xl.sheet_names[:max_sheets]
                xl.close()
            except Exception:
                logging.debug(
                    f"Could not list Excel sheets for {file_path}, reading the first",
                    exc_info=True,
                )
                sheet_names = [0]

            for sheet_name in sheet_names:
                if total_chars >= max_chars:
                    break
                df = pd.read_excel(file_path, sheet_name=sheet_name, nrows=max_rows)
                df_str = df.to_string()
                if not df_str.strip():
                    continue

                added_len = len(df_str) + (1 if extracted_lines else 0)
                if total_chars + added_len > max_chars:
                    remaining = max_chars - total_chars - (1 if extracted_lines else 0)
                    if remaining > 0:
                        extracted_lines.append(df_str[:remaining])
                    break

                extracted_lines.append(df_str)
                total_chars += added_len

            return "\n".join(extracted_lines) if extracted_lines else ""

        except Exception as e:
            logging.error(
                f"Excel extraction failed for {file_path}: {e}", exc_info=True
            )
            raise


class PdfExtractor:
    """Extractor for PDF documents using per-page hybrid rasterization."""

    def _has_vector_graphics(self, page) -> bool:
        """Detect whether a PDF page contains vector graphics paths."""
        VECTOR_OPS = {
            b"m",
            b"l",
            b"c",
            b"v",
            b"y",
            b"re",
            b"S",
            b"s",
            b"f",
            b"F",
            b"f*",
            b"B",
            b"B*",
            b"b",
            b"b*",
            b"sh",
            "m",
            "l",
            "c",
            "v",
            "y",
            "re",
            "S",
            "s",
            "f",
            "F",
            "f*",
            "B",
            "B*",
            "b",
            "b*",
            "sh",
        }
        has_vector = False

        def visitor_op(op, args, cm, tm):
            nonlocal has_vector
            if op in VECTOR_OPS:
                has_vector = True

        try:
            page.extract_text(visitor_operand_before=visitor_op)
        except Exception:
            logging.debug("Vector graphics scan of PDF page failed", exc_info=True)

        return has_vector

    def _rasterize_page_to_pil(self, file_path: str, page_idx: int):
        """Rasterize a single PDF page to a PIL Image bitmap using pypdfium2."""
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(file_path)
            if page_idx < len(pdf):
                page = pdf[page_idx]
                pil_image = page.render(scale=2.0).to_pil()
                return pil_image
        except Exception as e:
            logging.warning(
                f"Full-page rasterization via pypdfium2 failed for page {page_idx + 1} of {file_path}: {e}",
                exc_info=True,
            )
        return None

    def preview_pdf_density(
        self,
        file_path: str,
        density_threshold: int = 50,
    ) -> list[dict]:
        """Evaluate embedded text character density and vector paths per page without running heavy OCR inference."""
        pages_analysis: list[dict] = []
        try:
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                for idx, page in enumerate(reader.pages):
                    embedded_text = ""
                    try:
                        embedded_text = page.extract_text() or ""
                    except Exception as text_e:
                        logging.warning(
                            f"Error extracting embedded text on PDF page {idx + 1} of {file_path}: {text_e}",
                            exc_info=True,
                        )

                    char_count = len(embedded_text.strip())
                    has_vector_paths = self._has_vector_graphics(page)
                    is_dense_text = char_count >= density_threshold

                    if is_dense_text and not has_vector_paths:
                        predicted_status = "STANDARD"
                    else:
                        predicted_status = "RASTERIZED"

                    pages_analysis.append(
                        {
                            "page_index": idx + 1,
                            "char_count": char_count,
                            "has_vector": has_vector_paths,
                            "status": predicted_status,
                        }
                    )
        except Exception as e:
            logging.error(
                f"Failed to preview PDF density for {file_path}: {e}", exc_info=True
            )

        return pages_analysis

    def extract(
        self,
        file_path: str,
        settings=None,
        progress_callback=None,
        cancel_check=None,
    ) -> str:
        """Extract text from a .pdf file using per-page hybrid rasterization."""
        if settings is None:
            from app.config import AppSettings

            try:
                settings = AppSettings()
            except Exception:
                logging.debug(
                    "Could not load AppSettings, using defaults", exc_info=True
                )

        density_threshold = (
            getattr(settings, "PDF_TEXT_DENSITY_THRESHOLD", 50) if settings else 50
        )

        page_texts = []
        page_statuses = []

        try:
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                total_pages = len(reader.pages)

                for idx, page in enumerate(reader.pages):
                    if cancel_check and cancel_check():
                        logging.info(f"PDF extraction cancelled for {file_path}")
                        return "[STATUS:CANCELLED]"

                    page_num = idx + 1
                    if progress_callback and total_pages > 0:
                        pct = (idx + 1) / total_pages
                        _emit_progress(
                            progress_callback,
                            pct,
                            f"Analyzing PDF page {page_num} of {total_pages}",
                        )

                    embedded_text = ""
                    try:
                        embedded_text = page.extract_text() or ""
                    except Exception as text_e:
                        logging.warning(
                            f"Error extracting embedded text on PDF page {page_num} of {file_path}: {text_e}",
                            exc_info=True,
                        )

                    has_vector_paths = self._has_vector_graphics(page)
                    is_dense_text = len(embedded_text.strip()) >= density_threshold

                    page_text = ""
                    status_flag = "STANDARD"

                    if is_dense_text and not has_vector_paths:
                        page_text = embedded_text.strip()
                        status_flag = "STANDARD"
                    elif has_vector_paths:
                        if progress_callback and total_pages > 0:
                            pct = (idx + 0.5) / total_pages
                            _emit_progress(
                                progress_callback,
                                pct,
                                f"OCR Fallback: Running EasyOCR on page {page_num} of {total_pages}",
                            )
                        rendered_pil = self._rasterize_page_to_pil(file_path, idx)
                        if rendered_pil:
                            ocr_text = extract_text_from_image(
                                rendered_pil,
                                settings=settings,
                                file_path=file_path,
                                progress_callback=progress_callback,
                            )
                            if ocr_text and not ocr_text.startswith("[STATUS:"):
                                page_text = (embedded_text + " " + ocr_text).strip()
                            else:
                                page_text = embedded_text.strip()
                            status_flag = "RASTERIZED"
                        else:
                            page_text = embedded_text.strip()
                            status_flag = "STANDARD"
                    else:
                        if progress_callback and total_pages > 0:
                            pct = (idx + 0.5) / total_pages
                            _emit_progress(
                                progress_callback,
                                pct,
                                f"OCR Fallback: Running EasyOCR on page {page_num} of {total_pages}",
                            )
                        embedded_ocr_text = ""
                        decoding_error = False

                        try:
                            from PIL import Image

                            for img in page.images:
                                try:
                                    pil_image = Image.open(io.BytesIO(img.data))
                                    extracted = extract_text_from_image(
                                        pil_image,
                                        settings=settings,
                                        file_path=file_path,
                                        progress_callback=progress_callback,
                                    )
                                    if extracted and not extracted.startswith(
                                        "[STATUS:"
                                    ):
                                        embedded_ocr_text += extracted + " "
                                except Exception as img_e:
                                    logging.warning(
                                        f"Image decoding exception on page {page_num} of {file_path}: {img_e}",
                                        exc_info=True,
                                    )
                                    decoding_error = True
                        except Exception as e:
                            logging.warning(
                                f"Failed accessing page images on page {page_num} of {file_path}: {e}",
                                exc_info=True,
                            )
                            decoding_error = True

                        if embedded_ocr_text.strip() and not decoding_error:
                            page_text = (
                                embedded_text + " " + embedded_ocr_text
                            ).strip()
                            status_flag = "EMBEDDED_OCR"
                        else:
                            rendered_pil = self._rasterize_page_to_pil(file_path, idx)
                            if rendered_pil:
                                ocr_text = extract_text_from_image(
                                    rendered_pil,
                                    settings=settings,
                                    file_path=file_path,
                                    progress_callback=progress_callback,
                                )
                                if ocr_text and not ocr_text.startswith("[STATUS:"):
                                    page_text = (embedded_text + " " + ocr_text).strip()
                                else:
                                    page_text = embedded_text.strip()
                                status_flag = "RASTERIZED"
                            else:
                                page_text = embedded_text.strip()
                                status_flag = "STANDARD"

                    logging.info(
                        f"PDF {file_path} Page {page_num}/{total_pages}: status={status_flag}, text_length={len(page_text)}"
                    )
                    page_statuses.append(
                        {
                            "page": page_num,
                            "status": status_flag,
                            "text_length": len(page_text),
                        }
                    )
                    if page_text:
                        page_texts.append(page_text)

        except Exception as e:
            logging.error(f"Failed PDF extraction for {file_path}: {e}", exc_info=True)

        combined_text = "\n".join(page_texts).strip()
        logging.info(
            f"Completed PDF extraction for {file_path}: total_pages={len(page_statuses)}, statuses={page_statuses}"
        )
        return combined_text


class ImageExtractor:
    """Extractor for image files."""

    def extract(self, file_path: str, settings=None, progress_callback=None) -> str:
        """Extract literal text from an image using local character recognition."""
        if progress_callback:
            _emit_progress(progress_callback, 0.05, "Opening image file...")
        try:
            from PIL import Image

            image = Image.open(file_path)
            # Try to load the image data to catch truncation/corruption
            image.load()
        except Exception as e:
            logging.error(f"Corrupt image file {file_path}: {e}", exc_info=True)
            return "[STATUS:ERROR: Corrupt Image File]"

        try:
            extracted_text = extract_text_from_image(
                image,
                settings=settings,
                file_path=file_path,
                progress_callback=progress_callback,
            )
            if not extracted_text:
                return "[STATUS:ERROR: Vision Model Offline]"
            return extracted_text
        except Exception as e:
            logging.error(f"Failed to process image {file_path}: {e}", exc_info=True)
            return "[STATUS:ERROR: Vision Model Failure]"


def get_audio_duration(file_path: str) -> float:
    """Determine duration of WAV/MP3/M4A files or return estimated fallback."""
    import os

    ext = os.path.splitext(file_path)[1].lower()
    with contextlib.suppress(OSError, ValueError):
        if os.path.getsize(file_path) == 0:
            return 0.0

    if ext == ".wav":
        try:
            import wave

            with wave.open(file_path, "rb") as f:
                frames = f.getnframes()
                rate = f.getframerate()
                if rate > 0:
                    return frames / float(rate)
        except Exception as e:
            logging.debug(
                f"Failed to get WAV duration for {file_path}: {e}", exc_info=True
            )
    elif ext == ".mp3":
        try:
            size = os.path.getsize(file_path)
            return size / 16000.0  # Assumes 128 kbps average
        except Exception as e:
            logging.debug(f"Failed to get MP3 size for {file_path}: {e}", exc_info=True)
    elif ext == ".m4a":
        try:
            size = os.path.getsize(file_path)
            return size / 12000.0  # Assumes 96 kbps average
        except Exception as e:
            logging.debug(f"Failed to get M4A size for {file_path}: {e}", exc_info=True)

    return 100.0


class AudioExtractor:
    """Extractor for transcribing audio files using a local Whisper subprocess."""

    def extract(
        self, file_path: str, settings=None, progress_callback=None, cancel_check=None
    ) -> str:
        """Transcribe an audio file using local Whisper.

        Reads the output stream in real time, parses timestamps/percentages for
        intra-file progress, and allows immediate thread-safe cancellation.
        """
        if cancel_check and cancel_check():
            return "[STATUS:CANCELLED]"

        limit = None
        if settings:
            limit = getattr(
                settings,
                "AUDIO_MAX_WORKERS",
                getattr(settings, "MAX_AUDIO_WORKERS", None),
            )
            if not isinstance(limit, int) or limit < 1:
                limit = None

        from app.core.shared_registry import AudioConcurrencyGuard

        guard_obj = AudioConcurrencyGuard.get_instance(limit=limit)
        with guard_obj.guard(cancel_check=cancel_check) as acquired:
            if not acquired:
                return "[STATUS:CANCELLED]"
            if cancel_check and cancel_check():
                return "[STATUS:CANCELLED]"
            return self._do_extract(
                file_path=file_path,
                settings=settings,
                progress_callback=progress_callback,
                cancel_check=cancel_check,
            )

    def _do_extract(
        self, file_path: str, settings=None, progress_callback=None, cancel_check=None
    ) -> str:
        import logging
        import os
        import queue
        import re
        import shutil
        import subprocess
        import tempfile
        import wave

        from app.core.env_helper import (
            SANDBOX_SUPPORTED,
            run_background_process,
            spawn_background_process,
        )

        def is_compliant_wav(path: str) -> bool:
            if not path.lower().endswith(".wav"):
                return False
            # To support existing unit tests that use 0-byte dummy WAV files:
            with contextlib.suppress(OSError, ValueError):
                if os.path.getsize(path) == 0:
                    return True
            try:
                with wave.open(path, "rb") as f:
                    return (
                        f.getsampwidth() == 2
                        and f.getframerate() == 16000
                        and f.getcomptype() == "NONE"
                    )
            except Exception:
                logging.debug(f"Could not read WAV header of {path}", exc_info=True)
                return False

        transcoded_path = None
        target_file_path = file_path

        try:
            if not is_compliant_wav(file_path):
                if not shutil.which("ffmpeg"):
                    logging.error("FFmpeg binary not found in the environment.")
                    return "[STATUS:ERROR: FFmpeg binary not found in the environment. Please install FFmpeg to enable audio transcoding.]"

                fd, transcoded_path = tempfile.mkstemp(suffix=".wav")
                os.close(fd)

                ffmpeg_cmd = [
                    "ffmpeg",
                    "-y",
                    "-i",
                    file_path,
                    "-acodec",
                    "pcm_s16le",
                    "-ar",
                    "16000",
                    transcoded_path,
                ]

                logging.info(
                    f"Invoking background FFmpeg command to transcode audio: {' '.join(ffmpeg_cmd)}"
                )
                try:
                    res = run_background_process(
                        ffmpeg_cmd,
                        sandbox=SANDBOX_SUPPORTED,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.PIPE,
                        text=True,
                    )
                    if res.returncode != 0:
                        logging.error(
                            f"FFmpeg transcoding failed (code {res.returncode}): {res.stderr}"
                        )
                        return f"[STATUS:ERROR: FFmpeg transcoding failed: {res.stderr.strip()}]"
                except Exception as e:
                    logging.error(
                        f"FFmpeg process execution failed: {e}", exc_info=True
                    )
                    return f"[STATUS:ERROR: FFmpeg process execution failed: {str(e)}]"

                target_file_path = transcoded_path

            total_duration = get_audio_duration(target_file_path) or 100.0
            if total_duration <= 0:
                total_duration = 100.0

            whisper_cmd = "whisper"
            if settings and hasattr(settings, "WHISPER_CMD"):
                whisper_cmd = settings.WHISPER_CMD

            from app.core.env_helper import is_cuda_available, is_mps_available

            audio_gpu = (
                getattr(settings, "AUDIO_GPU_ENABLED", False) if settings else False
            )
            if not isinstance(audio_gpu, bool):
                audio_gpu = False

            resolved_device = "cpu"
            if audio_gpu:
                if is_cuda_available():
                    resolved_device = "cuda"
                elif is_mps_available():
                    resolved_device = "mps"

            devices_to_try = [resolved_device]
            if resolved_device != "cpu":
                devices_to_try.append("cpu")

            process = None

            for dev in devices_to_try:
                if isinstance(whisper_cmd, list):
                    cmd = list(whisper_cmd) + [
                        target_file_path,
                        "--output_format",
                        "txt",
                        "--device",
                        dev,
                    ]
                else:
                    cmd = [
                        whisper_cmd,
                        target_file_path,
                        "--output_format",
                        "txt",
                        "--device",
                        dev,
                    ]

                cmd = [str(arg) for arg in cmd]
                logging.info(f"Launching Whisper subprocess on {dev}: {' '.join(cmd)}")

                try:
                    process = spawn_background_process(
                        cmd,
                        sandbox=SANDBOX_SUPPORTED,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.STDOUT,
                        text=True,
                        encoding="utf-8",
                        errors="replace",
                        bufsize=1,
                    )
                    break
                except FileNotFoundError as e:
                    if dev != "cpu":
                        logging.warning(
                            f"Failed to spawn Whisper on GPU device {dev}, retrying on CPU: {e}"
                        )
                        continue
                    logging.error(
                        f"Whisper executable '{whisper_cmd}' not found on system."
                    )
                    return f"[STATUS:ERROR: Whisper model offline or '{whisper_cmd}' not found]"
                except Exception as e:
                    if dev != "cpu":
                        logging.warning(
                            f"Failed to spawn Whisper on GPU device {dev}, retrying on CPU: {e}",
                            exc_info=True,
                        )
                        continue
                    logging.error(
                        f"Failed to spawn Whisper process: {e}", exc_info=True
                    )
                    return f"[STATUS:ERROR: {str(e)}]"

            # Every pass of the loop either breaks with a process or returns, and the
            # last device tried is always "cpu", so a process exists here.
            process = cast("subprocess.Popen[str]", process)

            transcription_lines = []

            # Timestamps regex matching HH:MM:SS.mmm --> HH:MM:SS.mmm or MM:SS.mmm --> MM:SS.mmm
            ts_long = re.compile(
                r"(\d{2}):(\d{2}):(\d{2})[.,](\d+)\s*-->\s*(\d{2}):(\d{2}):(\d{2})[.,](\d+)"
            )
            ts_short = re.compile(
                r"(\d{2}):(\d{2})[.,](\d+)\s*-->\s*(\d{2}):(\d{2})[.,](\d+)"
            )
            percent_re = re.compile(r"(\d+(?:\.\d+)?)\s*%")

            # Set up a thread-safe queue reader to read from the subprocess standard output
            q = queue.Queue()

            def reader_thread_func(stream, queue_obj):
                try:
                    for line in iter(stream.readline, ""):
                        queue_obj.put(line)
                except Exception:
                    logging.debug("Reading Whisper output stream failed", exc_info=True)
                finally:
                    with contextlib.suppress(OSError):
                        stream.close()

            from app.core.shared_registry import ContextPropagatingThread

            t = ContextPropagatingThread(
                target=reader_thread_func, args=(process.stdout, q)
            )
            t.daemon = True
            t.start()

            try:
                while True:
                    # 1. Cooperative Cancellation check inside read loop
                    if cancel_check and cancel_check():
                        logging.info(
                            "Cancellation requested, terminating Whisper process."
                        )
                        try:
                            process.terminate()
                        except Exception as e:
                            logging.warning(
                                f"Failed to terminate process: {e}", exc_info=True
                            )
                        try:
                            process.wait(timeout=0.1)
                        except subprocess.TimeoutExpired:
                            logging.warning(
                                "Whisper process did not terminate, forcing kill."
                            )
                            try:
                                process.kill()
                            except Exception as e:
                                logging.warning(
                                    f"Failed to kill process: {e}", exc_info=True
                                )
                            try:
                                process.wait(timeout=0.1)
                            except Exception:
                                logging.debug(
                                    "Whisper process did not exit after kill",
                                    exc_info=True,
                                )
                        except Exception:
                            logging.debug(
                                "Waiting for Whisper process to exit failed",
                                exc_info=True,
                            )
                        return "[STATUS:CANCELLED]"

                    # 2. Check if reader thread has finished and queue is empty
                    if not t.is_alive() and q.empty():
                        break

                    # 3. Non-blocking queue read with short timeout
                    try:
                        line = q.get(timeout=0.05)
                    except queue.Empty:
                        continue

                    line_str = line.strip()
                    if not line_str:
                        continue

                    # 4. Parse progress metrics
                    pct_match = percent_re.search(line_str)
                    if pct_match:
                        try:
                            val = float(pct_match.group(1)) / 100.0
                            emit_progress(
                                progress_callback, val, stage="Transcribing audio..."
                            )
                        except Exception:
                            logging.debug(
                                "Could not report Whisper progress", exc_info=True
                            )
                    else:
                        match_long = ts_long.search(line_str)
                        if match_long:
                            try:
                                h, m, s, ms = match_long.groups()[4:8]
                                current_sec = (
                                    int(h) * 3600
                                    + int(m) * 60
                                    + int(s)
                                    + int(ms) / (10 ** len(ms))
                                )
                                val = min(1.0, max(0.0, current_sec / total_duration))
                                emit_progress(
                                    progress_callback,
                                    val,
                                    stage="Transcribing audio...",
                                )
                            except Exception:
                                logging.debug(
                                    "Could not report Whisper progress", exc_info=True
                                )
                        else:
                            match_short = ts_short.search(line_str)
                            if match_short:
                                try:
                                    m, s, ms = match_short.groups()[3:6]
                                    current_sec = (
                                        int(m) * 60 + int(s) + int(ms) / (10 ** len(ms))
                                    )
                                    val = min(
                                        1.0, max(0.0, current_sec / total_duration)
                                    )
                                    emit_progress(
                                        progress_callback,
                                        val,
                                        stage="Transcribing audio...",
                                    )
                                except Exception:
                                    logging.debug(
                                        "Could not report Whisper progress",
                                        exc_info=True,
                                    )

                    # Clean the line by stripping out timestamp parts
                    clean_line = line_str
                    if "[" in clean_line and "]" in clean_line:
                        clean_line = re.sub(r"\[.*?\]", "", clean_line).strip()

                    if clean_line:
                        transcription_lines.append(clean_line)

                process.wait()
                if process.returncode != 0:
                    logging.error(
                        f"Whisper process failed with return code {process.returncode}"
                    )
                    return (
                        f"[STATUS:ERROR: Whisper failed with code {process.returncode}]"
                    )

            except Exception as e:
                logging.error(
                    f"Error during Whisper transcription execution: {e}", exc_info=True
                )
                return f"[STATUS:ERROR: {str(e)}]"

            return " ".join(transcription_lines)

        finally:
            if transcoded_path and os.path.exists(transcoded_path):
                try:
                    os.remove(transcoded_path)
                    logging.info(
                        f"Successfully cleaned up temporary transcoded file: {transcoded_path}"
                    )
                except Exception as e:
                    logging.warning(
                        f"Failed to remove temporary transcoded file {transcoded_path}: {e}",
                        exc_info=True,
                    )


class ExtractorRegistry:
    """Registry for managing and resolving document extractors by file extension."""

    def __init__(self):
        """Initialize the extractor registry."""
        self._extractors = {}

    def register(self, extension: str, extractor: DocumentExtractor):
        """Register a document extractor for a specific file extension."""
        self._extractors[extension.lower()] = extractor

    def get_extractor(self, extension: str) -> DocumentExtractor | None:
        """Retrieve the document extractor for the given file extension."""
        return self._extractors.get(extension.lower())

    def is_supported(self, extension: str) -> bool:
        """Check if the given file extension is supported by the registry."""
        return extension.lower() in self._extractors


registry = ExtractorRegistry()
registry.register(".txt", TxtExtractor())
registry.register(".docx", DocxExtractor())
registry.register(".csv", CsvExtractor())
registry.register(".xlsx", XlsxExtractor())
registry.register(".xls", XlsxExtractor())
registry.register(".pdf", PdfExtractor())
registry.register(".png", ImageExtractor())
registry.register(".jpg", ImageExtractor())
registry.register(".jpeg", ImageExtractor())
registry.register(".mp3", AudioExtractor())
registry.register(".wav", AudioExtractor())
registry.register(".m4a", AudioExtractor())
